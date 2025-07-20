import PyPDF2
import openpyxl
import pandas as pd
import os
from typing import Optional, List
import pdfplumber
import pytesseract

def extract_text_from_pdf(file_path: str) -> str:
    """
    Extract text from a PDF file using OCR for image-based PDFs.
    
    Args:
        file_path: Path to the PDF file
        
    Returns:
        str: Extracted text from the PDF file
    """
    text = ""
    try:
        # First try standard text extraction
        with open(file_path, "rb") as file:
            reader = PyPDF2.PdfReader(file)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text.strip():
                    text += page_text
        
        # If no text found, use OCR
        if not text.strip():
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    img = page.to_image(resolution=150)
                    page_text = pytesseract.image_to_string(img.original)
                    text += page_text + "\n"
    except Exception as e:
        return f"Error reading PDF: {str(e)}"
    
    return text


def extract_data_from_excel(file_path: str) -> str:
    """
    Extract data from an Excel file and convert to text format.
    
    Args:
        file_path: Path to the Excel file
        
    Returns:
        str: Formatted text representation of Excel data
    """
    try:
        # Read all sheets from the Excel file
        excel_data = pd.read_excel(file_path, sheet_name=None, engine='openpyxl')
        
        combined_text = f"EXCEL FILE WITH {len(excel_data)} SHEETS:\n"
        
        for sheet_name, df in excel_data.items():
            combined_text += f"\n{'='*60}\n"
            combined_text += f"SHEET: {sheet_name.upper()}\n"
            combined_text += f"{'='*60}\n"
            
            # Add sheet summary
            combined_text += f"Rows: {len(df)}, Columns: {len(df.columns)}\n"
            combined_text += f"Column Headers: {', '.join(df.columns.astype(str))}\n\n"
            
            # Convert DataFrame to string with better formatting
            combined_text += "DATA:\n"
            combined_text += df.to_string(index=False, na_rep='', max_rows=100)
            combined_text += "\n\n"
        
        return combined_text
    except Exception as e:
        return f"Error reading Excel: {str(e)}"


def get_excel_info(file_path: str) -> dict:
    """
    Get metadata from an Excel file.
    
    Args:
        file_path: Path to the Excel file
        
    Returns:
        dict: Excel metadata
    """
    try:
        workbook = openpyxl.load_workbook(file_path, read_only=True)
        return {
            "sheet_names": workbook.sheetnames,
            "num_sheets": len(workbook.sheetnames),
            "file_size": os.path.getsize(file_path)
        }
    except Exception as e:
        return {"error": str(e)}


def get_pdf_info(file_path: str) -> dict:
    """
    Get metadata from a PDF file.
    
    Args:
        file_path: Path to the PDF file
        
    Returns:
        dict: PDF metadata
    """
    try:
        with open(file_path, "rb") as file:
            reader = PyPDF2.PdfReader(file)
            return {
                "num_pages": len(reader.pages),
                "metadata": reader.metadata,
                "is_encrypted": reader.is_encrypted
            }
    except Exception as e:
        return {"error": str(e)}
